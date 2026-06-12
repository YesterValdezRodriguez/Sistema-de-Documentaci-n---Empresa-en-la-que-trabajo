"""Exportación a Excel (.xlsx) con estilo corporativo FARACH."""
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

AZUL_FARACH = '1B3A6B'
GRIS_CLARO = 'F4F6F9'


def exportar_excel(titulo_reporte, nombre_hoja, encabezados, filas):
    """Genera un archivo Excel en memoria con el estilo corporativo.

    - titulo_reporte: título mostrado en la primera fila
    - nombre_hoja: nombre de la pestaña
    - encabezados: lista de strings
    - filas: lista de listas/tuplas con los valores
    Devuelve BytesIO listo para send_file.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = nombre_hoja[:31]

    borde = Border(bottom=Side(style='thin', color='CCCCCC'))
    num_cols = len(encabezados)

    # Fila de título con fecha de generación
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
    celda_titulo = ws.cell(row=1, column=1)
    celda_titulo.value = f'{titulo_reporte} — Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}'
    celda_titulo.font = Font(bold=True, size=13, color=AZUL_FARACH)
    celda_titulo.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height = 24

    # Encabezados
    for col, encabezado in enumerate(encabezados, start=1):
        celda = ws.cell(row=2, column=col, value=encabezado)
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill(start_color=AZUL_FARACH, end_color=AZUL_FARACH, fill_type='solid')
        celda.alignment = Alignment(horizontal='center', vertical='center')

    # Datos con filas alternas
    for idx, fila in enumerate(filas):
        num_fila = idx + 3
        for col, valor in enumerate(fila, start=1):
            celda = ws.cell(row=num_fila, column=col, value=valor)
            celda.border = borde
            if idx % 2 == 1:
                celda.fill = PatternFill(start_color=GRIS_CLARO, end_color=GRIS_CLARO,
                                         fill_type='solid')

    # Auto-ajuste de ancho de columnas
    for col in range(1, num_cols + 1):
        max_largo = len(str(encabezados[col - 1]))
        for fila in filas:
            if col <= len(fila) and fila[col - 1] is not None:
                max_largo = max(max_largo, len(str(fila[col - 1])))
        ws.column_dimensions[get_column_letter(col)].width = min(max_largo + 3, 60)

    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)
    return salida
